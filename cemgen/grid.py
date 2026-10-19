"""Загрузка файла матрицы в сетку строк (DESIGN.md, 9.1).

Сетка — строки файла с номерами; ячейки — строки без пробелов по краям, хвостовые пустые
ячейки отброшены. Пустые строки и комментарии (первая ячейка начинается с ``#``) в сетку не
входят, но номера строк остаются номерами файла или листа: сообщения проверки указывают на
настоящую строку.

* CSV: UTF-8 (BOM на входе допускается), разделитель ``;``, кавычки по RFC 4180, LF или CRLF.
* XLSX: лист ``CEM`` (если лист один — он); числа читаются как текст (целое ``2.0`` → ``2``,
  иначе кратчайшая запись с точкой), дата → ``ГГГГ-ММ-ДД``, формула → сохранённое значение;
  формула без значения и объединённые ячейки — ошибка CEM-E009.

SHA-256 сетки (:meth:`Grid.digest`) считается по нормализованной сетке: строки через ``\\n``,
ячейки через ``\\t``. Одна и та же матрица в CSV и в XLSX (``cemgen convert``) даёт одну сетку,
поэтому и одинаковые сгенерированные файлы.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from cemgen.diagnostics import Diagnostic, MatrixFileError

SHEET_NAME = "CEM"
CSV_DELIMITER = ";"
SUFFIXES = (".csv", ".xlsx")


@dataclass(frozen=True)
class GridRow:
    """Строка сетки: номер строки файла (с 1) и ячейки (колонка A — индекс 0)."""

    line: int
    cells: tuple[str, ...]

    def cell(self, col: int) -> str:
        """Ячейка по номеру колонки с единицы; за концом строки — пустая."""
        return self.cells[col - 1] if 0 < col <= len(self.cells) else ""

    @property
    def first(self) -> str:
        return self.cell(1)


@dataclass(frozen=True)
class Grid:
    """Содержательные строки матрицы (без пустых и комментариев)."""

    rows: tuple[GridRow, ...]

    def digest(self) -> str:
        """Первые 16 шестнадцатеричных цифр SHA-256 нормализованной сетки."""
        text = "\n".join("\t".join(row.cells) for row in self.rows)
        return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


@dataclass(frozen=True)
class RawRow:
    """Строка файла как есть (для ``convert``): номер и ячейки, включая комментарии и пустые."""

    line: int
    cells: tuple[str, ...]


def _normalize(cells: list[str]) -> tuple[str, ...]:
    stripped = [cell.strip() for cell in cells]
    while stripped and not stripped[-1]:
        stripped.pop()
    return tuple(stripped)


def _is_comment(cells: tuple[str, ...]) -> bool:
    return bool(cells) and cells[0].startswith("#")


def grid_from_raw(raw: list[RawRow]) -> Grid:
    """Сетка из строк файла: пустые строки и комментарии отбрасываются."""
    rows: list[GridRow] = []
    for raw_row in raw:
        cells = _normalize(list(raw_row.cells))
        if not cells or _is_comment(cells):
            continue
        rows.append(GridRow(raw_row.line, cells))
    return Grid(tuple(rows))


def _file_error(message: str, hint: str) -> MatrixFileError:
    return MatrixFileError(Diagnostic("CEM-E001", message, hint))


def read_csv_raw(path: Path) -> list[RawRow]:
    """Все строки CSV с номерами строк файла (запись в кавычках на нескольких строках — по первой)."""
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise _file_error(f"файл не читается: {exc.strerror or exc}", "проверьте путь и права доступа") from exc
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise _file_error(
            f"кодировка не UTF-8 (байт {exc.start + 1})", "сохраните CSV в кодировке UTF-8 (с BOM или без)"
        ) from exc
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=CSV_DELIMITER, quotechar='"', strict=True)
    rows: list[RawRow] = []
    line_end = 0
    try:
        for record in reader:
            rows.append(RawRow(line_end + 1, tuple(record)))
            line_end = reader.line_num
    except csv.Error as exc:
        raise _file_error(
            f"CSV не разбирается в строке {reader.line_num}: {exc}",
            "закройте кавычки ячейки или удвойте кавычку внутри неё (RFC 4180)",
        ) from exc
    return rows


def _xlsx_text(value: object) -> str:
    """Значение ячейки XLSX → текст по правилам 9.1."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        text = format(Decimal(repr(value)), "f")
        return text
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    return str(value)


def read_xlsx_raw(path: Path) -> list[RawRow]:
    """Строки листа ``CEM`` как текст; ошибки листа — CEM-E009."""
    import openpyxl
    from openpyxl.utils.cell import get_column_letter

    try:
        book_values = openpyxl.load_workbook(path, data_only=True)
        book_formulas = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:  # openpyxl бросает разные исключения на повреждённый файл
        raise _file_error(f"XLSX не открывается: {exc}", "сохраните книгу заново в формате .xlsx") from exc
    names = book_values.sheetnames
    if SHEET_NAME in names:
        name = SHEET_NAME
    elif len(names) == 1:
        name = names[0]
    else:
        raise MatrixFileError(
            Diagnostic(
                "CEM-E009",
                f"в книге {len(names)} листов и нет листа {SHEET_NAME}",
                f"назовите лист с матрицей {SHEET_NAME} или оставьте в книге один лист",
            )
        )
    sheet = book_values[name]
    sheet_formulas = book_formulas[name]
    merged = sorted(sheet.merged_cells.ranges, key=lambda rng: (rng.min_row, rng.min_col))
    if merged:
        first = merged[0]
        raise MatrixFileError(
            Diagnostic(
                "CEM-E009",
                f"объединённые ячейки {first.coord}",
                "разъедините ячейки: каждая ячейка сетки — отдельное значение",
                first.min_row,
                first.min_col,
            )
        )
    rows: list[RawRow] = []
    max_col = sheet.max_column
    for row_cells in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=max_col):
        texts: list[str] = []
        line = row_cells[0].row if row_cells else len(rows) + 1
        for cell in row_cells:
            value = cell.value
            if value is None and cell.row is not None and cell.column is not None:
                source = sheet_formulas.cell(row=cell.row, column=cell.column)
                if source.data_type == "f":
                    raise MatrixFileError(
                        Diagnostic(
                            "CEM-E009",
                            f"формула без сохранённого значения в {get_column_letter(cell.column)}{cell.row}",
                            "пересчитайте и сохраните книгу в Excel или замените формулу значением",
                            cell.row,
                            cell.column,
                        )
                    )
            texts.append(_xlsx_text(value))
        rows.append(RawRow(int(line or 0), tuple(texts)))
    return rows


def read_raw(path: Path) -> list[RawRow]:
    """Строки файла матрицы (CSV или XLSX) без нормализации."""
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise _file_error(f"расширение «{path.suffix}» не поддерживается", "сохраните матрицу как .csv или .xlsx")
    if not path.is_file():
        raise _file_error("файл не найден", "проверьте путь к матрице")
    if suffix == ".csv":
        return read_csv_raw(path)
    return read_xlsx_raw(path)


def read_grid(path: Path) -> Grid:
    """Сетка матрицы из CSV или XLSX; нечитаемый файл — :class:`MatrixFileError`."""
    return grid_from_raw(read_raw(path))
