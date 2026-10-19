"""Разбор значений ячеек (9.1, 9.3) и загрузка CSV/XLSX: сетка, номера строк, CEM-E009, CSV = XLSX."""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from cemgen.convert import convert
from cemgen.diagnostics import MatrixFileError, column_letter
from cemgen.grid import read_grid, read_raw
from cemgen.validate import load
from cemgen.values import (
    DurationError,
    format_decimal_ru,
    format_duration_ru,
    is_ascii_text,
    is_iso_date,
    is_tag,
    matrix_mark,
    parse_bool,
    parse_duration_ms,
    parse_number,
)
from pytests.conftest import DEMO_CSV


@pytest.mark.parametrize(
    ("text", "ms"),
    [
        ("500ms", 500),
        ("2s", 2000),
        ("1.5s", 1500),
        ("1,5s", 1500),
        ("15m", 900_000),
        ("8h", 28_800_000),
        ("10", 10_000),
        ("0", 0),
        ("2S", 2000),
        ("100MS", 100),
        ("0.1s", 100),
    ],
)
def test_duration_ok(text: str, ms: int) -> None:
    assert parse_duration_ms(text) == ms


@pytest.mark.parametrize("text", ["2 s", "abc", "1.0005s", "s", "2d", "", "1e3s"])
def test_duration_bad(text: str) -> None:
    with pytest.raises(DurationError):
        parse_duration_ms(text)


def test_duration_negative_is_returned_for_range_check() -> None:
    assert parse_duration_ms("-1s") == -1000


@pytest.mark.parametrize(
    ("text", "value"),
    [("yes", True), ("Да", True), ("1", True), ("TRUE", True), ("no", False), ("нет", False), ("0", False)],
)
def test_bool(text: str, value: bool) -> None:
    assert parse_bool(text) is value


def test_bool_bad() -> None:
    assert parse_bool("maybe") is None
    assert parse_bool("") is None


def test_numbers() -> None:
    assert parse_number("0.90") == Decimal("0.90")
    assert parse_number("0,9") == Decimal("0.9")
    assert parse_number("-5") == Decimal(-5)
    assert parse_number(".5") == Decimal("0.5")
    for bad in ("1 000", "1.000,5", "5V", "1e3", "", "--1"):
        assert parse_number(bad) is None


def test_tags() -> None:
    assert is_tag("PSHH101")
    assert is_tag("A")
    assert is_tag("A" + "1" * 23)
    for bad in ("A" + "1" * 24, "PSHH-101", "pshh101", "1ABC", "P_T", "P T", "ПТ101", ""):
        assert not is_tag(bad)


def test_marks_and_misc() -> None:
    assert matrix_mark("X") and matrix_mark("x") and matrix_mark("Х") and matrix_mark("х")
    assert matrix_mark("") is False
    assert matrix_mark("XX") is None
    assert is_iso_date("2026-10-16") and not is_iso_date("2026-02-30") and not is_iso_date("16.10.2026")
    assert is_ascii_text("V-101") and not is_ascii_text("В-101") and not is_ascii_text("a\tb")
    assert column_letter(1) == "A" and column_letter(26) == "Z" and column_letter(27) == "AA"
    assert column_letter(52) == "AZ" and column_letter(53) == "BA"
    assert format_decimal_ru(Decimal("0.90")) == "0,9"
    assert format_decimal_ru(Decimal("85")) == "85"
    assert format_duration_ru(899_900) == "14 мин 59,9 с"
    assert format_duration_ru(28_800_000) == "8 ч"
    assert format_duration_ru(500) == "500 мс"


def test_csv_line_numbers_skip_comments(tmp_path: Path) -> None:
    grid = read_grid(DEMO_CSV)
    first = grid.rows[0]
    assert first.first == "[header]"
    assert first.line == 6  # пять строк комментариев сверху
    assert all(not row.first.startswith("#") for row in grid.rows)


def test_csv_quoted_cell_with_semicolon(tmp_path: Path) -> None:
    path = tmp_path / "q.csv"
    path.write_text('a;"b;c";"d ""e"""\n\n# comment\nx\n', encoding="utf-8")
    raw = read_raw(path)
    assert raw[0].cells == ("a", "b;c", 'd "e"')
    grid = read_grid(path)
    assert [r.line for r in grid.rows] == [1, 4]


def _xlsx_from_rows(path: Path, rows: list[list[Any]], title: str = "CEM") -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = title
    for r, cells in enumerate(rows, 1):
        for c, value in enumerate(cells, 1):
            if value is not None:
                sheet.cell(row=r, column=c, value=value)
    book.save(path)
    return path


def test_xlsx_values_as_text(tmp_path: Path) -> None:
    path = _xlsx_from_rows(tmp_path / "v.xlsx", [[2.0, 0.9, 85, dt.date(2026, 10, 16), True, " txt "]])
    raw = read_raw(path)
    assert raw[0].cells == ("2", "0.9", "85", "2026-10-16", "true", " txt ")


def test_xlsx_sheet_selection(tmp_path: Path) -> None:
    path = tmp_path / "s.xlsx"
    book = openpyxl.Workbook()
    first = book.active
    assert first is not None
    first.title = "Notes"
    first["A1"] = "not the matrix"
    book.create_sheet("CEM")["A1"] = "[header]"
    book.save(path)
    assert read_raw(path)[0].cells == ("[header]",)


def test_xlsx_errors_e009(tmp_path: Path) -> None:
    two = tmp_path / "two.xlsx"
    book = openpyxl.Workbook()
    book.create_sheet("Other")
    book.save(two)
    with pytest.raises(MatrixFileError) as info:
        read_raw(two)
    assert info.value.diagnostic.code == "CEM-E009"

    merged = _xlsx_from_rows(tmp_path / "m.xlsx", [["a", "b"]])
    book = openpyxl.load_workbook(merged)
    book.active.merge_cells("A1:B1")  # type: ignore[union-attr]
    book.save(merged)
    with pytest.raises(MatrixFileError) as info:
        read_raw(merged)
    assert info.value.diagnostic.code == "CEM-E009"
    assert (info.value.diagnostic.row, info.value.diagnostic.col) == (1, 1)

    formula = _xlsx_from_rows(tmp_path / "f.xlsx", [["=1+1"]])
    with pytest.raises(MatrixFileError) as info:
        read_raw(formula)
    assert info.value.diagnostic.code == "CEM-E009"


def test_xlsx_broken_file_is_e001(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"not a zip")
    result = load(path)
    assert [d.code for d in result.diagnostics] == ["CEM-E001"]


def test_csv_and_xlsx_give_same_grid_and_model(tmp_path: Path) -> None:
    xlsx = tmp_path / "cem.xlsx"
    convert(DEMO_CSV, xlsx)
    csv_result = load(DEMO_CSV)
    xlsx_result = load(xlsx)
    assert xlsx_result.diagnostics == ()
    assert csv_result.matrix == xlsx_result.matrix
    back = tmp_path / "back.csv"
    convert(xlsx, back)
    assert read_grid(back) == read_grid(DEMO_CSV)
    # номера строк сохраняются: сообщения указывают на ту же строку в обоих форматах
    assert [r.line for r in read_grid(xlsx).rows] == [r.line for r in read_grid(DEMO_CSV).rows]


def test_convert_rejects_unknown_target(tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        convert(DEMO_CSV, tmp_path / "x.ods")


def test_committed_xlsx_matches_csv() -> None:
    """``examples/separator/cem.xlsx`` (``cemgen convert``) — та же матрица, что CSV (10.2)."""
    xlsx = DEMO_CSV.with_suffix(".xlsx")
    if not xlsx.is_file():
        pytest.skip(
            "cem.xlsx ещё не создан: python -m cemgen convert examples/separator/cem.csv examples/separator/cem.xlsx"
        )
    assert load(xlsx).matrix == load(DEMO_CSV).matrix
